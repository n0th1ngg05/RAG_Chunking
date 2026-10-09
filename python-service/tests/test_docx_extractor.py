import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from app.services.extractor import extract_docx


class DocxExtractorTests(unittest.TestCase):
    def test_extract_docx_keeps_paragraphs_with_missing_styles(self):
        fake_doc = SimpleNamespace(
            paragraphs=[
                SimpleNamespace(text="Styled heading", style=SimpleNamespace(name="Heading 1")),
                SimpleNamespace(text="Paragraph with missing style object", style=None),
                SimpleNamespace(text="Paragraph with missing style name", style=SimpleNamespace(name=None)),
                SimpleNamespace(text="Styled list item", style=SimpleNamespace(name="List Bullet")),
            ]
        )

        with patch("app.services.extractor.DocxDocument", return_value=fake_doc):
            pages = extract_docx("missing-style.docx")

        self.assertEqual(len(pages), 1)
        self.assertIn("Styled heading", pages[0][1])
        self.assertIn("Paragraph with missing style object", pages[0][1])
        self.assertIn("Paragraph with missing style name", pages[0][1])
        self.assertIn("- Styled list item", pages[0][1])

    def test_extract_docx_preserves_normal_styled_heading_paragraph_and_list(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "styled.docx"
            doc = Document()
            doc.add_heading("Project Overview", level=1)
            doc.add_paragraph("This paragraph should remain plain text.")
            doc.add_paragraph("This item should become a list marker.", style="List Bullet")
            doc.save(path)

            pages = extract_docx(str(path))

        self.assertEqual(len(pages), 1)
        text = pages[0][1]
        self.assertIn("Project Overview", text)
        self.assertIn("This paragraph should remain plain text.", text)
        self.assertIn("- This item should become a list marker.", text)


if __name__ == "__main__":
    unittest.main()
